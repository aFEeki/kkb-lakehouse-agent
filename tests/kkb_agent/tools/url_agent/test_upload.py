import hashlib
from datetime import UTC, datetime
from io import BytesIO

import pytest
from openpyxl import Workbook

from kkb_agent.tools.url_agent import (
    UPLOAD_NOTE,
    DocumentKind,
    UnsupportedContentTypeError,
    UploadError,
    UploadTooLargeError,
    create_content_type_router,
    read_upload,
)

UPLOADED_AT = datetime(2026, 9, 15, 10, 30, tzinfo=UTC)

HTML = b"<!doctype html><html><body><p>Altin islemleri</p></body></html>"


def upload(filename, body, **changes):
    fields = {
        "router": create_content_type_router(),
        "uploaded_by": "ege",
        "uploaded_at": UPLOADED_AT,
    }
    return read_upload(filename, body, **(fields | changes))


def xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Veriler"
    sheet.append(["Tarih", "Fiyat"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_an_uploaded_file_goes_through_the_same_parsers_as_a_fetch():
    result = upload("sayfa.html", HTML)

    assert result.document.kind is DocumentKind.HTML
    assert result.document.extraction_method == "beautifulsoup-lxml"
    assert "Altin islemleri" in result.document.text


def test_the_document_itself_says_it_was_supplied_by_a_person():
    result = upload("sayfa.html", HTML)
    assert result.document.notes[0] == UPLOAD_NOTE


def test_provenance_records_who_supplied_it_when_and_which_bytes():
    result = upload("kmp_au.pdf", HTML, declared_content_type="text/html")

    provenance = result.provenance
    assert provenance.filename == "kmp_au.pdf"
    assert provenance.uploaded_by == "ege"
    assert provenance.uploaded_at == UPLOADED_AT
    assert provenance.byte_count == len(HTML)
    assert provenance.content_sha256 == hashlib.sha256(HTML).hexdigest()
    assert provenance.declared_content_type == "text/html"


def test_the_bytes_decide_the_type_not_the_filename():
    """A person naming a file .pdf does not make it one, and the reverse happens too:
    a saved page arrives as report.pdf while actually being HTML."""

    assert upload("report.pdf", HTML).document.kind is DocumentKind.HTML
    assert upload("veriler.txt", xlsx_bytes()).document.kind is DocumentKind.EXCEL


def test_an_html_table_saved_as_xls_is_read_rather_than_rejected():
    """Observed in the wild: plenty of systems export an HTML table and name it .xls,
    because Excel opens it. Trusting the extension hands those bytes to xlrd, which
    refuses them outright ("Expected BOF record") and the rows are simply lost."""

    exported = (
        b"<html><head><meta charset='utf-8'></head><body><table>"
        b"<tr><td>Tarih</td><td>Tutar</td></tr>"
        b"<tr><td>2026-09-15</td><td>1.234,56</td></tr>"
        b"</table></body></html>"
    )

    result = upload("agreement-list.xls", exported)

    assert result.document.kind is DocumentKind.HTML
    assert result.document.tables[0].rows == (
        ("Tarih", "Tutar"),
        ("2026-09-15", "1.234,56"),
    )


def test_a_file_whose_bytes_are_a_broken_pdf_fails_as_a_pdf():
    """Routing still goes by the bytes, so the failure names the real problem rather
    than silently treating a truncated PDF as text."""
    from kkb_agent.tools.url_agent import ContentExtractionError

    with pytest.raises(ContentExtractionError, match="malformed"):
        upload("rapor.txt", b"%PDF-1.4\ntruncated")


def test_excel_and_text_uploads_are_accepted():
    workbook = upload("veriler.xlsx", xlsx_bytes())
    assert workbook.document.kind is DocumentKind.EXCEL
    assert workbook.document.tables[0].caption == "Veriler"

    text = upload("not.txt", "Altın 4.512,30".encode(), declared_content_type="text/plain")
    assert text.document.kind is DocumentKind.TEXT
    assert "4.512,30" in text.document.text


def test_an_image_upload_needs_the_same_ocr_backend_a_fetch_would():
    class FakeOCR:
        def read_page_images(self, images):
            return tuple("Resimden okunan" for _ in images)

    result = upload(
        "tablo.png",
        b"\x89PNG\r\n\x1a\npayload",
        router=create_content_type_router(ocr=FakeOCR()),
    )

    assert result.document.kind is DocumentKind.IMAGE
    assert result.document.text == "Resimden okunan"
    assert result.document.notes[0] == UPLOAD_NOTE


def test_an_unreadable_upload_is_refused_rather_than_guessed_at():
    with pytest.raises(UnsupportedContentTypeError):
        upload("mystery.bin", b"\x00\x01\x02\x03")


def test_an_empty_upload_is_refused():
    with pytest.raises(UploadError, match="empty"):
        upload("bos.pdf", b"")


def test_an_oversized_upload_is_refused_with_both_numbers():
    with pytest.raises(UploadTooLargeError, match="over the limit"):
        upload("buyuk.html", HTML, max_bytes=10)


def test_a_non_bytes_body_is_refused():
    with pytest.raises(UploadError, match="must be bytes"):
        upload("sayfa.html", "not bytes")


def test_upload_provenance_is_immutable():
    result = upload("sayfa.html", HTML)
    with pytest.raises(ValueError):
        result.provenance.uploaded_by = "someone else"
